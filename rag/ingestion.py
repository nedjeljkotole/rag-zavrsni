"""Sloj za unos podataka: učitavanje, normalizacija i segmentacija dokumenata."""
from __future__ import annotations

import hashlib
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable, List

from .models import Chunk, Document

SUPPORTED = {".txt", ".md", ".html", ".htm", ".pdf", ".docx"}


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: List[str] = []
        self._skip = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip = True
        if tag in ("p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = False

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def read_text(path: Path) -> str:
    """Izdvaja tekst iz datoteke. PDF i DOCX zahtijevaju opcionalne pakete."""
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in (".html", ".htm"):
        p = _HTMLText()
        p.feed(path.read_text(encoding="utf-8", errors="replace"))
        return "".join(p.parts)
    if suffix == ".pdf":
        from pypdf import PdfReader  # pip install pypdf
        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        import docx  # pip install python-docx
        return "\n\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    raise ValueError(f"Nepodržan format: {path.suffix}")


def normalize_text(text: str) -> str:
    """Konzervativna normalizacija: ujednačuje razmake i prazne retke, sadržaj ne mijenja."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def load_sources(folder: str | Path) -> List[Document]:
    """Učitava sve podržane datoteke iz mape kao objekte Document."""
    folder = Path(folder)
    docs: List[Document] = []
    for path in sorted(folder.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            text = normalize_text(read_text(path))
            if not text:
                continue
            title = text.splitlines()[0].strip("# ").strip()[:120]
            docs.append(Document(document_id=path.stem, title=title, text=text,
                                 source=str(path.relative_to(folder))))
    return docs


_SENT_SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-ZČĆĐŠŽ0-9„\"(])")


def _words(s: str) -> int:
    return len(s.split())


def split_text(text: str, max_words: int = 200, overlap_words: int = 30) -> List[str]:
    """Dijeli tekst po odlomcima i rečenicama uz gornju granicu duljine i preklapanje.

    - kratki susjedni odlomci spajaju se dok ne dosegnu max_words,
    - predugi odlomak dijeli se po rečenicama,
    - svaki sljedeći segment počinje s posljednjih ~overlap_words riječi prethodnog.
    """
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    units: List[str] = []
    for p in paragraphs:
        if _words(p) <= max_words:
            units.append(p)
        else:
            for s in (s.strip() for s in _SENT_SPLIT.split(p) if s.strip()):
                if _words(s) <= max_words:
                    units.append(s)
                else:  # predugi odlomak bez prepoznatljivih rečenica dijeli se po riječima
                    w = s.split()
                    units.extend(" ".join(w[i:i + max_words]) for i in range(0, len(w), max_words))

    chunks: List[str] = []
    current: List[str] = []
    for u in units:
        if current and _words(" ".join(current + [u])) > max_words:
            chunks.append("\n".join(current))
            tail = " ".join(current).split()[-overlap_words:] if overlap_words else []
            current = [" ".join(tail)] if tail else []
        current.append(u)
    if current:
        chunks.append("\n".join(current))
    return chunks


def content_hash(text: str) -> str:
    norm = re.sub(r"\s+", " ", text.strip().lower())
    return hashlib.sha256(norm.encode("utf-8")).hexdigest()[:16]


def split_and_attach_metadata(documents: Iterable[Document], max_words: int = 200,
                              overlap_words: int = 30, min_words: int = 3) -> List[Chunk]:
    """Segmentira dokumente i svakom segmentu dodaje metapodatke.

    Prazni, prekratki i duplicirani segmenti (isti sažetak sadržaja) se preskaču.
    """
    seen = set()
    out: List[Chunk] = []
    for doc in documents:
        for i, part in enumerate(split_text(doc.text, max_words, overlap_words)):
            if _words(part) < min_words:
                continue
            h = content_hash(part)
            if h in seen:
                continue
            seen.add(h)
            out.append(Chunk(chunk_id=f"{doc.document_id}#{i}", document_id=doc.document_id,
                             position=i, text=part, title=doc.title, source=doc.source,
                             version=doc.version, content_hash=h))
    return out
